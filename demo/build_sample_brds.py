"""Render the Markdown sample BRDs to DOCX and PDF (real formats for exercising the parsers).

    cd backend && uv run python ../demo/build_sample_brds.py
"""

import html
import re
from pathlib import Path

import pymupdf
from docx import Document

HERE = Path(__file__).parent / "brds"


def md_blocks(text: str) -> list[tuple[str, object]]:
    """Tiny Markdown subset: #-headings, | tables |, - lists, paragraphs."""
    out: list[tuple[str, object]] = []
    lines = text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
        elif line.startswith("#"):
            m = re.match(r"(#+)\s+(.*)", line)
            out.append(("h", (len(m.group(1)), m.group(2))))
            i += 1
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.fullmatch(r"\|[\s:|-]+\|", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            out.append(("table", rows))
        elif line.startswith("- "):
            items = []
            while i < len(lines) and lines[i].startswith("- "):
                items.append(lines[i][2:].strip())
                i += 1
            out.append(("ul", items))
        else:
            out.append(("p", line.strip()))
            i += 1
    return out


def to_docx(blocks: list[tuple[str, object]], path: Path) -> None:
    doc = Document()
    for kind, value in blocks:
        if kind == "h":
            level, title = value  # type: ignore[misc]
            doc.add_heading(title, level=min(level, 4))
        elif kind == "p":
            doc.add_paragraph(value)
        elif kind == "ul":
            for item in value:  # type: ignore[attr-defined]
                doc.add_paragraph(item, style="List Bullet")
        elif kind == "table":
            rows = value  # type: ignore[assignment]
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    t.cell(r, c).text = cell
    doc.save(path)


def to_pdf(blocks: list[tuple[str, object]], path: Path) -> None:
    parts = [
        '<html><head><style>table{border-collapse:collapse;} '
        'th,td{border:1px solid #000000; padding:4px;} th{background-color:#dddddd;}</style></head>'
        '<body style="font-family: helvetica; font-size: 10pt;">'
    ]
    sizes = {1: "20pt", 2: "15pt", 3: "12.5pt", 4: "11pt"}
    for kind, value in blocks:
        if kind == "h":
            level, title = value  # type: ignore[misc]
            parts.append(f'<h{level} style="font-size:{sizes.get(level, "11pt")};">{html.escape(title)}</h{level}>')
        elif kind == "p":
            parts.append(f"<p>{html.escape(value)}</p>")  # type: ignore[arg-type]
        elif kind == "ul":
            parts.append("<ul>" + "".join(f"<li>{html.escape(i)}</li>" for i in value) + "</ul>")  # type: ignore[attr-defined]
        elif kind == "table":
            rows = value  # type: ignore[assignment]
            parts.append('<table width="100%">')
            for r, row in enumerate(rows):
                tag = "th" if r == 0 else "td"
                parts.append("<tr>" + "".join(f"<{tag}>{html.escape(c)}</{tag}>" for c in row) + "</tr>")
            parts.append("</table><p></p>")
    parts.append("</body></html>")
    story = pymupdf.Story(html="".join(parts))
    writer = pymupdf.DocumentWriter(str(path))
    mediabox = pymupdf.paper_rect("a4")
    where = mediabox + (50, 60, -50, -60)
    more = True
    while more:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()


def main() -> None:
    for md in sorted(HERE.glob("*.md")):
        blocks = md_blocks(md.read_text())
        to_docx(blocks, md.with_suffix(".docx"))
        to_pdf(blocks, md.with_suffix(".pdf"))
        print("built", md.stem, [p.suffix for p in (md.with_suffix(".docx"), md.with_suffix(".pdf"))])


if __name__ == "__main__":
    main()
