# ruff: noqa: F811, ARG001
"""Review (edit, versions, bulk, regenerate), the RTM and the exports."""

import io
import uuid
from collections.abc import Callable
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlmodel import Session, select

from app.core.config import settings
from app.core.db import engine
from app.qa.llm.fake import RuleChatModel
from tests.qa.test_generation import (  # noqa: F401  (fixtures are imported on purpose)
    World,
    author_model,
    cases,
    generate,
    world,
)

API = settings.API_V1_STR
H = dict[str, str]


def by_title(client: TestClient, h: H, w: World, title: str) -> dict[str, Any]:
    return next(c for c in cases(client, h, w) if c["title"] == title)


def rtm(client: TestClient, h: H, w: World) -> dict[str, Any]:
    r = client.get(f"{API}/projects/{w.project_id}/rtm", headers=h)
    assert r.status_code == 200, r.text
    return r.json()  # type: ignore[no-any-return]


# ------------------------------------------------------------------------------------ editing
def test_editing_a_case_versions_it_and_requires_approval_again(
    world: Callable[[], World], client: TestClient, superuser_token_headers: H
) -> None:
    h, w = superuser_token_headers, world()
    generate(client, h, w)
    c = by_title(client, h, w, "Title is required on 'Add Item'")
    client.patch(
        f"{API}/tests/{c['id']}/status", headers=h, json={"status": "approved"}
    )
    r = client.patch(
        f"{API}/tests/{c['id']}",
        headers=h,
        json={
            "title": "Title must be filled in",
            "expected_result": "An error is shown",
            "priority": "P1",
            "note": "clearer",
        },
    )
    assert r.status_code == 200, r.text
    e = r.json()
    assert e["status"] == "edited" and e["version"] == 2 and e["priority"] == "P1"
    assert e["tc_key"] == c["tc_key"]
    versions = client.get(f"{API}/tests/{c['id']}/versions", headers=h).json()
    assert [v["version"] for v in versions] == [2, 1]
    assert (
        versions[0]["change_type"] == "edited"
        and versions[0]["change_note"] == "clearer"
    )
    assert versions[1]["content"]["title"] == "Title is required on 'Add Item'"
    # restore the generated wording: a new version, not a rewrite of history
    back = client.post(
        f"{API}/tests/{c['id']}/restore", headers=h, json={"version": 1}
    ).json()
    assert back["title"] == "Title is required on 'Add Item'" and back["version"] == 3
    assert len(client.get(f"{API}/tests/{c['id']}/versions", headers=h).json()) == 3


def test_edits_are_validated_and_hand_edited_steps_are_flagged(
    world: Callable[[], World], client: TestClient, superuser_token_headers: H
) -> None:
    h, w = superuser_token_headers, world()
    generate(client, h, w)
    c = cases(client, h, w)[0]
    assert (
        client.patch(
            f"{API}/tests/{c['id']}", headers=h, json={"priority": "P9"}
        ).status_code
        == 422
    )
    assert (
        client.patch(
            f"{API}/tests/{c['id']}", headers=h, json={"title": ""}
        ).status_code
        == 422
    )
    r = client.patch(
        f"{API}/tests/{c['id']}",
        headers=h,
        json={"rendered_steps": ["1. Do it by hand"]},
    )
    assert r.json()["evidence"]["hand_edited"] is True and r.json()[
        "rendered_steps"
    ] == ["1. Do it by hand"]


def test_bulk_review_and_ownership(
    world: Callable[[], World],
    client: TestClient,
    superuser_token_headers: H,
    normal_user_token_headers: H,
) -> None:
    h, w = superuser_token_headers, world()
    generate(client, h, w)
    ids = [c["id"] for c in cases(client, h, w)[:5]]
    r = client.post(
        f"{API}/projects/{w.project_id}/tests/bulk",
        headers=h,
        json={"ids": ids, "action": "approve"},
    )
    assert r.json() == {"updated": 5}
    assert len(cases(client, h, w, status="approved")) == 5
    assert (
        client.post(
            f"{API}/projects/{w.project_id}/tests/bulk",
            headers=normal_user_token_headers,
            json={"ids": ids, "action": "reject"},
        ).status_code
        == 403
    )
    assert (
        client.patch(
            f"{API}/tests/{ids[0]}",
            headers=normal_user_token_headers,
            json={"title": "x"},
        ).status_code
        == 403
    )
    assert (
        client.get(
            f"{API}/tests/{ids[0]}/versions", headers=normal_user_token_headers
        ).status_code
        == 403
    )


# -------------------------------------------------------------------------------- regeneration
def test_regeneration_with_feedback_reaches_the_author_and_keeps_human_work(
    world: Callable[[], World],
    client: TestClient,
    superuser_token_headers: H,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    h, w = superuser_token_headers, world()
    generate(client, h, w)
    scenario = next(
        c for c in cases(client, h, w, origin="llm") if c["req_key"] == "FR-1"
    )
    edited = by_title(client, h, w, "Title is required on 'Add Item'")
    client.patch(
        f"{API}/tests/{edited['id']}",
        headers=h,
        json={"expected_result": "A reviewer wrote this."},
    )
    prompts: list[str] = []

    def spy(prompt: str) -> dict[str, Any]:
        prompts.append(prompt)
        return author_model(prompt)

    monkeypatch.setattr(
        "app.qa.runs.kinds.generate.get_chat_model",
        lambda *a, **k: RuleChatModel(handler=spy, tool_name="AuthorResult"),
    )
    r = client.post(
        f"{API}/tests/{scenario['id']}/regenerate",
        headers=h,
        json={"feedback": "Cover the empty list case"},
    )
    assert r.status_code == 200 and r.json()["kind"] == "regenerate"
    from tests.qa.test_brd_api import run_worker

    assert run_worker(r.json()["id"]) == "succeeded"
    assert (
        len(prompts) == 1 and "Requirement FR-1" in prompts[0]
    )  # only that requirement was redone
    assert "Cover the empty list case" in prompts[0]
    after = {c["id"]: c for c in cases(client, h, w)}
    assert (
        after[edited["id"]]["expected_result"] == "A reviewer wrote this."
    )  # human edit survives
    assert scenario["id"] not in after  # the unreviewed scenario was replaced
    new = [c for c in after.values() if c["origin"] == "llm" and c["req_key"] == "FR-1"]
    assert new and new[0]["tc_key"] != scenario["tc_key"]
    v = client.get(f"{API}/tests/{new[0]['id']}/versions", headers=h).json()
    assert v[0]["change_type"] == "regenerated" and "empty list" in v[0]["change_note"]


def test_changing_a_requirement_queues_drafts_and_marks_approved_cases(
    world: Callable[[], World], client: TestClient, superuser_token_headers: H
) -> None:
    h, w = superuser_token_headers, world()
    generate(client, h, w)
    own = [
        c
        for c in cases(client, h, w)
        if c["req_key"] == "FR-2" and c["origin"] == "llm"
    ]
    keep = by_title(client, h, w, "Title is required on 'Add Item'")  # belongs to FR-1
    client.patch(
        f"{API}/tests/{keep['id']}/status", headers=h, json={"status": "approved"}
    )
    assert own
    r = client.patch(
        f"{API}/requirements/{w.reqs['FR-1']}",
        headers=h,
        json={"description": "Items can be added with a title and a body."},
    )
    assert r.status_code == 200
    after = {c["id"]: c for c in cases(client, h, w)}
    assert (
        after[keep["id"]]["status"] == "approved"
        and after[keep["id"]]["evidence"]["requirement_changed"] is True
    )
    stale = [
        c
        for c in after.values()
        if c["req_key"] == "FR-1" and c["status"] == "needs_regen"
    ]
    assert stale and all(c["evidence"]["requirement_changed"] for c in stale)
    assert all(
        c["status"] != "needs_regen" for c in after.values() if c["req_key"] == "FR-2"
    )


# ------------------------------------------------------------------------------------- RTM
def test_rtm_reconciles_with_the_test_table_and_explains_every_gap(
    world: Callable[[], World], client: TestClient, superuser_token_headers: H
) -> None:
    h, w = superuser_token_headers, world()
    generate(client, h, w)
    all_cases = cases(client, h, w)
    client.patch(
        f"{API}/tests/{all_cases[0]['id']}/status",
        headers=h,
        json={"status": "rejected"},
    )
    m = rtm(client, h, w)
    live = [c for c in cases(client, h, w) if c["status"] != "rejected"]
    assert m["test_cases"] == len(
        live
    )  # RTM counts == table counts (rejected excluded)
    rows = {r["req_key"]: r for r in m["rows"]}
    assert sum(r["tests_own"] for r in m["rows"]) == len(live)
    assert rows["FR-1"]["coverage"] == "covered" and rows["FR-1"]["mapped_screens"] == 1
    assert (
        rows["FR-5"]["tests_linked"] > 0
    )  # identical cases kept under FR-1 also trace to FR-5
    assert (
        rows["FR-3"]["coverage"] == "uncovered"
        and "no matching screen" in rows["FR-3"]["reason"]
    )
    assert rows["FR-4"]["needs_decision"] and "wording" in rows["FR-4"]["reason"]
    assert "non-functional" in rows["NFR-1"]["reason"]
    assert m["covered"] + m["uncovered"] == m["requirements"] == 6
    assert sum(rows["FR-1"]["by_type"].values()) == rows["FR-1"]["tests_own"]


# --------------------------------------------------------------------------------- exports
def test_test_case_exports_follow_the_template_and_neutralise_formulas(
    world: Callable[[], World], client: TestClient, superuser_token_headers: H
) -> None:
    h, w = superuser_token_headers, world()
    generate(client, h, w)
    c = by_title(client, h, w, "Title is required on 'Add Item'")
    client.patch(
        f"{API}/tests/{c['id']}",
        headers=h,
        json={"title": '=HYPERLINK("http://evil","x")'},
    )
    base = f"{API}/projects/{w.project_id}"
    csv_r = client.get(f"{base}/exports/tests", headers=h, params={"format": "csv"})
    assert (
        csv_r.status_code == 200
        and "attachment" in csv_r.headers["content-disposition"]
    )
    lines = csv_r.content.decode("utf-8-sig").splitlines()
    assert lines[0].startswith("Test Case ID,Requirement ID,Module,Title")
    assert (
        "'=HYPERLINK" in csv_r.text and ",=HYPERLINK" not in csv_r.text
    )  # a formula never starts a cell
    # a template: fewer columns, renamed, reordered, another ID prefix
    t = client.put(
        f"{base}/template",
        headers=h,
        json={
            "id_prefix": "ACME",
            "columns": [
                {"key": "title", "header": "Scenario"},
                {"key": "tc_key", "header": "Ref"},
                {"key": "expectation_basis", "header": ""},
            ],
        },
    )
    assert (
        t.status_code == 200
        and t.json()["columns"][2]["header"] == "Expected result basis"
    )
    js = client.get(
        f"{base}/exports/tests", headers=h, params={"format": "json"}
    ).json()
    assert list(js[0]) == ["Scenario", "Ref", "Expected result basis"]
    assert js[0]["Ref"].startswith("ACME-")
    xl = client.get(f"{base}/exports/tests", headers=h, params={"format": "xlsx"})
    ws = load_workbook(io.BytesIO(xl.content)).active
    assert [c.value for c in ws[1]] == ["Scenario", "Ref", "Expected result basis"]
    assert ws.max_row - 1 == len(
        [x for x in cases(client, h, w) if x["status"] != "rejected"]
    )


def test_template_rejects_unknown_and_repeated_columns(
    world: Callable[[], World], client: TestClient, superuser_token_headers: H
) -> None:
    h, w = superuser_token_headers, world()
    url = f"{API}/projects/{w.project_id}/template"
    assert (
        client.put(
            url, headers=h, json={"columns": [{"key": "nope", "header": ""}]}
        ).status_code
        == 422
    )
    assert (
        client.put(
            url,
            headers=h,
            json={
                "columns": [
                    {"key": "title", "header": ""},
                    {"key": "title", "header": "again"},
                ]
            },
        ).status_code
        == 422
    )
    assert (
        client.put(
            url,
            headers=h,
            json={"id_prefix": "1bad", "columns": [{"key": "title", "header": ""}]},
        ).status_code
        == 422
    )
    assert (
        len(client.get(url, headers=h).json()["columns"]) == 11
    )  # untouched: the default template


def test_rtm_gap_and_knowledge_model_exports(
    world: Callable[[], World], client: TestClient, superuser_token_headers: H
) -> None:
    h, w = superuser_token_headers, world()
    generate(client, h, w)
    base = f"{API}/projects/{w.project_id}/exports"
    x = load_workbook(
        io.BytesIO(
            client.get(f"{base}/rtm", headers=h, params={"format": "xlsx"}).content
        )
    )
    ws = x.active
    assert (
        ws[1][0].value == "Requirement ID" and ws.max_row == 7
    )  # header + 6 requirements
    rtm_csv = client.get(
        f"{base}/rtm", headers=h, params={"format": "csv"}
    ).content.decode("utf-8-sig")
    assert "FR-3" in rtm_csv and "no matching screen" in rtm_csv
    km = load_workbook(
        io.BytesIO(
            client.get(f"{base}/km", headers=h, params={"format": "xlsx"}).content
        )
    )
    assert km.sheetnames == ["Screens", "Fields"]
    fields = [r[2].value for r in km["Fields"].iter_rows(min_row=2)]
    assert "Title" in fields and "Quantity" in fields
    assert (
        client.get(f"{base}/gaps", headers=h, params={"format": "json"}).status_code
        == 200
    )
    assert client.get(f"{base}/nope", headers=h).status_code == 404
    assert (
        client.get(f"{base}/tests", headers=h, params={"format": "pdf"}).status_code
        == 422
    )


def test_exports_are_private_to_the_project_owner(
    world: Callable[[], World],
    client: TestClient,
    superuser_token_headers: H,
    normal_user_token_headers: H,
) -> None:
    w = world()
    for path in ("rtm", "exports/tests", "template", "metrics"):
        r = client.get(
            f"{API}/projects/{w.project_id}/{path}", headers=normal_user_token_headers
        )
        assert r.status_code == 403, path


def test_metrics_report_time_and_tokens_per_kind(
    world: Callable[[], World], client: TestClient, superuser_token_headers: H
) -> None:
    h, w = superuser_token_headers, world()
    generate(client, h, w)
    from app.qa.models import Run

    with Session(
        engine
    ) as sess:  # the fake model reports no usage: set what a real run records
        run = sess.exec(
            select(Run).where(
                Run.project_id == uuid.UUID(w.project_id), Run.kind == "generate"
            )
        ).one()
        run.summary = {**run.summary, "input_tokens": 1200, "output_tokens": 300}
        sess.add(run)
        sess.commit()
    m = client.get(f"{API}/projects/{w.project_id}/metrics", headers=h).json()
    gen = next(k for k in m["by_kind"] if k["kind"] == "generate")
    assert (
        gen["runs"] == 1
        and (gen["input_tokens"], gen["output_tokens"]) == (1200, 300)
        and gen["seconds"] >= 0
    )
    assert m["input_tokens"] == 1200 and m["runs"][0]["kind"] == "generate"


def test_a_case_from_before_versioning_gets_a_baseline_and_unchanged_text_is_not_an_edit(
    world: Callable[[], World], client: TestClient, superuser_token_headers: H
) -> None:
    h, w = superuser_token_headers, world()
    generate(client, h, w)
    c = cases(client, h, w)[0]
    with Session(engine) as s:  # simulate a case generated before M8
        from app.qa.models import TestCaseVersion

        for v in s.exec(
            select(TestCaseVersion).where(
                TestCaseVersion.test_case_id == uuid.UUID(c["id"])
            )
        ):
            s.delete(v)
        s.commit()
    same = client.patch(
        f"{API}/tests/{c['id']}",
        headers=h,
        json={"rendered_steps": c["rendered_steps"], "title": c["title"]},
    )
    assert (
        same.json()["version"] == 1 and same.json()["status"] == c["status"]
    )  # nothing differed
    r = client.patch(
        f"{API}/tests/{c['id']}",
        headers=h,
        json={"rendered_steps": c["rendered_steps"], "title": "New wording"},
    )
    assert r.json()["version"] == 2 and "hand_edited" not in r.json()["evidence"]
    vs = client.get(f"{API}/tests/{c['id']}/versions", headers=h).json()
    assert [v["version"] for v in vs] == [2, 1] and vs[1]["content"]["title"] == c[
        "title"
    ]
