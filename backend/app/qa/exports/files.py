"""Serialise rows as CSV, JSON or Excel. Cells that a spreadsheet would run as a formula are neutralised."""

import csv
import io
import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MEDIA = {
    "csv": "text/csv; charset=utf-8",
    "json": "application/json",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
MAX_CELL = 32000  # Excel's limit is 32,767 characters


def safe(value: Any) -> Any:
    """A BRD- or app-derived string that starts with = + - @ would execute in Excel/LibreOffice."""
    if isinstance(value, str):
        if value.startswith(FORMULA_START):
            value = "'" + value
        return value[:MAX_CELL]
    return value


def as_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "\n".join(str(v) for v in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def to_csv(headers: list[str], rows: list[list[Any]]) -> bytes:
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow([safe(h) for h in headers])
    for r in rows:
        w.writerow([safe(as_text(c)) for c in r])
    return ("﻿" + out.getvalue()).encode()  # BOM: Excel opens UTF-8 correctly


def to_json(headers: list[str], rows: list[list[Any]]) -> bytes:
    return json.dumps(
        [dict(zip(headers, r, strict=True)) for r in rows],
        ensure_ascii=False,
        indent=2,
        default=str,
    ).encode()


def to_xlsx(sheets: dict[str, tuple[list[str], list[list[Any]]]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]
    for name, (headers, rows) in sheets.items():
        ws = wb.create_sheet(name[:31])
        ws.append([safe(h) for h in headers])
        for r in rows:
            ws.append([safe(as_text(c)) for c in r])
        fill = PatternFill("solid", fgColor="1F3A5F")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        for i, h in enumerate(headers, 1):
            longest = max(
                [len(str(h))]
                + [len(as_text(r[i - 1]).split("\n")[0]) for r in rows[:200]]
            )
            ws.column_dimensions[get_column_letter(i)].width = min(
                max(10, longest + 2), 60
            )
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        if rows:
            ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render(
    fmt: str, sheets: dict[str, tuple[list[str], list[list[Any]]]]
) -> tuple[bytes, str]:
    """`sheets` has one entry for csv/json (the first is used) and any number for xlsx."""
    name = next(iter(sheets))
    headers, rows = sheets[name]
    if fmt == "csv":
        return to_csv(headers, rows), MEDIA["csv"]
    if fmt == "json":
        return to_json(headers, rows), MEDIA["json"]
    return to_xlsx(sheets), MEDIA["xlsx"]
